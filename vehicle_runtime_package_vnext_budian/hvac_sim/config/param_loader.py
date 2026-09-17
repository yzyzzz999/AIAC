"""Load initial calibration parameter tables from xlsx / csv / json."""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass, field, fields, replace
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union

from hvac_sim.afe.params import AFEParams
from hvac_sim.chtd.params import CHTDParams

from .param_mapping import apply_record_with_aliases

VALID_TARGET_TYPES = frozenset({"airflow", "chtd", "pmv", "ekf", "geometry"})

_CHTD_CANONICAL_TO_DEPRECATED_ALIAS = {
    "CHTD_Dt_P": "Dt",
    "CHTD_QgainCo_P": "QgainCo",
    "CHTD_QlossCo_P": "QlossCo",
    "CHTD_CHTDSdlTi_P": "SdlTi",
    "CHTD_AirCpAtb_P": "AirCpAtb",
    "CHTD_HeadAirVAtb_P": "HeadAirVAtb",
    "CHTD_FeetAirVAtb_P": "FeetAirVAtb",
}

_NAME_ALIASES = frozenset(
    {"name", "param", "parameter", "python_name", "simulink_var", "var", "变量名"}
)
_VALUE_ALIASES = frozenset({"value", "val", "default", "initial", "数值"})
_UNIT_ALIASES = frozenset({"unit", "units", "docunits", "单位"})
_SOURCE_ALIASES = frozenset({"source", "origin", "来源"})
_TARGET_ALIASES = frozenset({"target_type", "target", "domain", "类别"})
_NOTE_ALIASES = frozenset({"note", "notes", "description", "remark", "备注"})


@dataclass
class InitialParamRecord:
    name: str
    value: Any
    unit: str = ""
    source: str = ""
    target_type: str = ""
    note: str = ""


@dataclass
class InitialParamsDocument:
    schema: str = "initial_calibration_params_v1"
    classification: str = "initial_guess_not_final_calibration"
    source_files: List[str] = field(default_factory=list)
    parameters: List[InitialParamRecord] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "schema": self.schema,
            "classification": self.classification,
            "source_files": list(self.source_files),
            "parameters": [asdict(item) for item in self.parameters],
        }


class ParamTableError(ValueError):
    """Invalid parameter table input."""


def infer_target_type(name: str) -> str:
    key = str(name).strip()
    upper = key.upper()
    if upper.startswith("AFE_FH") or upper.startswith("AFE_") or upper.endswith("FLAPRESSCO") or "RESSCO" in upper:
        return "airflow"
    if upper.startswith("CHTD_"):
        return "chtd"
    if any(token in upper for token in ("EKF", "PROCESS_VAR", "MEASUREMENT_VAR", "SURFACE_TO_AIR")):
        return "ekf"
    if any(
        token in key.lower()
        for token in (
            "area",
            "distance",
            "distribution",
            "attenuation",
            "geometry",
            "outlet",
        )
    ):
        return "geometry"
    if any(token in upper for token in ("PMV", "CLO", "MET", "BIAS")):
        return "pmv"
    if upper.startswith(("F1_", "F2_", "S_", "FAN_")):
        return "airflow"
    return "airflow"


def _normalize_header(value: str) -> str:
    return re.sub(r"[\s_\-]+", "", str(value).strip().lower())


def _pick_column(fieldnames: Sequence[str], aliases: frozenset[str]) -> Optional[str]:
    normalized_aliases = {_normalize_header(alias) for alias in aliases}
    for name in fieldnames:
        if _normalize_header(name) in normalized_aliases:
            return name
    return None


def _coerce_value(raw: Any) -> Any:
    if raw is None:
        return None
    if isinstance(raw, (int, float, list)):
        return raw
    text = str(raw).strip()
    if not text:
        return None
    if text.startswith("[") and text.endswith("]"):
        try:
            return json.loads(text.replace("'", '"'))
        except json.JSONDecodeError:
            pass
    try:
        if "." in text or "e" in text.lower():
            return float(text)
        return int(text)
    except ValueError:
        return text


def _record_from_mapping(row: Mapping[str, Any], *, default_source: str) -> Optional[InitialParamRecord]:
    name_raw = row.get("name") or row.get("python_name") or row.get("simulink_var")
    if name_raw is None:
        return None
    name = str(name_raw).strip()
    if not name:
        return None
    value = row.get("value")
    if value is None and "Value" in row:
        value = row.get("Value")
    value = _coerce_value(value)
    if value is None:
        return None
    unit = str(row.get("unit") or row.get("DocUnits") or "").strip()
    source = str(row.get("source") or default_source).strip()
    target_type = str(row.get("target_type") or infer_target_type(name)).strip().lower()
    if target_type not in VALID_TARGET_TYPES:
        target_type = infer_target_type(name)
    note = str(row.get("note") or row.get("description") or "").strip()
    if row.get("is_placeholder"):
        note = (note + " placeholder").strip()
    return InitialParamRecord(
        name=name,
        value=value,
        unit=unit,
        source=source,
        target_type=target_type,
        note=note,
    )


def _parse_csv_rows(path: Path) -> List[InitialParamRecord]:
    records: List[InitialParamRecord] = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            return records
        name_col = _pick_column(reader.fieldnames, _NAME_ALIASES)
        value_col = _pick_column(reader.fieldnames, _VALUE_ALIASES)
        if name_col is None or value_col is None:
            raise ParamTableError(f"csv {path} must include name and value columns")
        unit_col = _pick_column(reader.fieldnames, _UNIT_ALIASES)
        source_col = _pick_column(reader.fieldnames, _SOURCE_ALIASES)
        target_col = _pick_column(reader.fieldnames, _TARGET_ALIASES)
        note_col = _pick_column(reader.fieldnames, _NOTE_ALIASES)
        for row in reader:
            payload = {
                "name": row.get(name_col),
                "value": row.get(value_col),
                "unit": row.get(unit_col) if unit_col else "",
                "source": row.get(source_col) if source_col else path.name,
                "target_type": row.get(target_col) if target_col else "",
                "note": row.get(note_col) if note_col else "",
            }
            rec = _record_from_mapping(payload, default_source=path.name)
            if rec is not None:
                records.append(rec)
    return records


def _parse_json_document(path: Path) -> InitialParamsDocument:
    doc = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(doc, dict):
        raise ParamTableError(f"json {path} root must be an object")

    if isinstance(doc.get("parameters"), list):
        params = []
        for item in doc["parameters"]:
            if not isinstance(item, dict):
                continue
            rec = _record_from_mapping(item, default_source=str(item.get("source") or path.name))
            if rec is not None:
                params.append(rec)
        return InitialParamsDocument(
            schema=str(doc.get("schema", "initial_calibration_params_v1")),
            classification=str(
                doc.get("classification", "initial_guess_not_final_calibration")
            ),
            source_files=[str(path)],
            parameters=params,
        )

    if isinstance(doc.get("params"), list):
        params = []
        for item in doc["params"]:
            if not isinstance(item, dict):
                continue
            name = item.get("python_name") or item.get("simulink_var")
            rec = _record_from_mapping(
                {
                    "name": name,
                    "value": item.get("value"),
                    "unit": item.get("unit", ""),
                    "source": item.get("source", path.name),
                    "target_type": infer_target_type(str(name or "")),
                    "note": item.get("note", ""),
                },
                default_source=path.name,
            )
            if rec is not None:
                if item.get("is_placeholder"):
                    rec.note = (rec.note + " initial guess placeholder").strip()
                params.append(rec)
        return InitialParamsDocument(source_files=[str(path)], parameters=params)

    raise ParamTableError(f"json {path} missing 'parameters' or 'params' list")


def _parse_xlsx_rows(path: Path) -> List[InitialParamRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParamTableError(
            "reading .xlsx requires openpyxl; install openpyxl or provide csv/json"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    sheet = wb["base"] if "base" in sheet_names else wb[sheet_names[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    fieldnames = [str(cell).strip() if cell is not None else "" for cell in header]
    name_col_idx = None
    value_col_idx = None
    for idx, name in enumerate(fieldnames):
        norm = _normalize_header(name)
        if norm in _NAME_ALIASES and name_col_idx is None:
            name_col_idx = idx
        if norm in _VALUE_ALIASES and value_col_idx is None:
            value_col_idx = idx
    if name_col_idx is None or value_col_idx is None:
        name_col_idx = 0
        value_col_idx = 1 if len(fieldnames) > 1 else None
    if value_col_idx is None:
        raise ParamTableError(f"xlsx {path} must include value column")

    unit_idx = next((i for i, n in enumerate(fieldnames) if _normalize_header(n) in _UNIT_ALIASES), None)
    source_idx = next((i for i, n in enumerate(fieldnames) if _normalize_header(n) in _SOURCE_ALIASES), None)
    target_idx = next((i for i, n in enumerate(fieldnames) if _normalize_header(n) in _TARGET_ALIASES), None)
    note_idx = next((i for i, n in enumerate(fieldnames) if _normalize_header(n) in _NOTE_ALIASES), None)

    records: List[InitialParamRecord] = []
    for row in rows:
        if row is None or name_col_idx >= len(row):
            continue
        payload = {
            "name": row[name_col_idx],
            "value": row[value_col_idx] if value_col_idx < len(row) else None,
            "unit": row[unit_idx] if unit_idx is not None and unit_idx < len(row) else "",
            "source": row[source_idx] if source_idx is not None and source_idx < len(row) else path.name,
            "target_type": row[target_idx] if target_idx is not None and target_idx < len(row) else "",
            "note": row[note_idx] if note_idx is not None and note_idx < len(row) else "",
        }
        rec = _record_from_mapping(payload, default_source=path.name)
        if rec is not None:
            records.append(rec)
    return records


def load_initial_params(path: Union[str, Path]) -> InitialParamsDocument:
    """Load parameter records from xlsx, csv, or json."""
    src = Path(path)
    if not src.exists():
        raise ParamTableError(f"parameter table not found: {src}")

    suffix = src.suffix.lower()
    if suffix == ".json":
        return _parse_json_document(src)
    if suffix == ".csv":
        return InitialParamsDocument(source_files=[str(src)], parameters=_parse_csv_rows(src))
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return InitialParamsDocument(source_files=[str(src)], parameters=_parse_xlsx_rows(src))
    raise ParamTableError(f"unsupported parameter table format: {suffix}")


def write_initial_calibration_params(
    document: InitialParamsDocument,
    output_path: Union[str, Path],
    *,
    pretty: bool = True,
) -> Path:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(document.to_dict(), ensure_ascii=False, indent=2 if pretty else None)
    out.write_text(text + "\n", encoding="utf-8")
    return out


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _rel_source(path: Path) -> str:
    root = _repo_root()
    try:
        return str(path.relative_to(root)).replace("\\", "/")
    except ValueError:
        return str(path)


def build_default_initial_params_document() -> InitialParamsDocument:
    """Build initial-guess document from bundled repo exports when xlsx is absent."""
    root = _repo_root()
    sources: List[str] = []
    parameters: List[InitialParamRecord] = []

    afe_json = root / "simulink_conversion_package" / "params" / "afe_params.json"
    if afe_json.exists():
        doc = _parse_json_document(afe_json)
        sources.append(_rel_source(afe_json))
        for item in doc.parameters:
            parameters.append(
                InitialParamRecord(
                    name=item.name,
                    value=item.value,
                    unit=item.unit,
                    source=item.source or "param.xlsx",
                    target_type="airflow",
                    note=(
                        "Initial guess from other-vehicle param.xlsx export; "
                        "not final calibration for current vehicle"
                    ),
                )
            )

    for extra in (
        InitialParamRecord(
            name="OsaFlapRessCo",
            value=300.0,
            unit="Pa*s^2/m^6",
            source="placeholder",
            target_type="airflow",
            note="Missing in param.xlsx; placeholder until bench calibration",
        ),
        InitialParamRecord(
            name="RecFlapRessCo",
            value=300.0,
            unit="Pa*s^2/m^6",
            source="placeholder",
            target_type="airflow",
            note="Missing in param.xlsx; placeholder until bench calibration",
        ),
    ):
        parameters.append(extra)

    fan_json = root / "simulink_conversion_package" / "params" / "fan_params.json"
    if fan_json.exists():
        fan_doc = json.loads(fan_json.read_text(encoding="utf-8"))
        sources.append(_rel_source(fan_json))
        for fan_name, payload in fan_doc.get("fans", {}).items():
            for key in ("n0", "p0", "a0"):
                if key in payload:
                    parameters.append(
                        InitialParamRecord(
                            name=f"{fan_name}_{key}",
                            value=float(payload[key]),
                            unit="",
                            source="VFC_Config.sldd",
                            target_type="airflow",
                            note="Fan curve placeholder; requires bench calibration",
                        )
                    )

    geom_json = Path(__file__).with_name("geometry_defaults.json")
    if geom_json.exists():
        geom = json.loads(geom_json.read_text(encoding="utf-8"))
        sources.append(_rel_source(geom_json))
        for outlet_name, outlet in geom.get("outlets", {}).items():
            for field_name, unit in (
                ("effective_area_m2", "m2"),
                ("distance_to_head_m", "m"),
                ("distribution_ratio", "-"),
            ):
                if field_name in outlet:
                    parameters.append(
                        InitialParamRecord(
                            name=f"{outlet_name}.{field_name}",
                            value=float(outlet[field_name]),
                            unit=unit,
                            source="geometry_defaults.json",
                            target_type="geometry",
                            note=str(outlet.get("note", "measured=false placeholder")),
                        )
                    )
        att = geom.get("global", {}).get("attenuation_k", {})
        if att:
            parameters.append(
                InitialParamRecord(
                    name="attenuation_k",
                    value=float(att.get("value", 1.0)),
                    unit="1/m2",
                    source="geometry_defaults.json",
                    target_type="geometry",
                    note="Distance attenuation placeholder; not formal calibration",
                )
            )

    chtd_json = root / "simulink_conversion_package" / "params" / "chtd_params.json"
    if chtd_json.exists():
        chtd = json.loads(chtd_json.read_text(encoding="utf-8"))
        sources.append(_rel_source(chtd_json))
        for name, payload in chtd.get("global_scalars", {}).items():
            if isinstance(payload, dict) and "value" in payload:
                parameters.append(
                    InitialParamRecord(
                        name=name,
                        value=payload["value"],
                        unit=str(payload.get("unit", "")),
                        source="ParamList.xlsx",
                        target_type="chtd",
                        note="CHTD placeholder scalar from other-vehicle param list",
                    )
                )

    return InitialParamsDocument(source_files=sources, parameters=parameters)


def export_default_initial_calibration_params(
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    if output_path is None:
        output_path = (
            _repo_root()
            / "simulink_conversion_package"
            / "python_targets"
            / "initial_calibration_params.json"
        )
    doc = build_default_initial_params_document()
    return write_initial_calibration_params(doc, output_path)


def bundled_initial_params_path() -> Path:
    return (
        _repo_root()
        / "simulink_conversion_package"
        / "python_targets"
        / "initial_calibration_params.json"
    )


def resolve_initial_params(
    path: Optional[Union[str, Path]] = None,
) -> Tuple[InitialParamsDocument, str]:
    """Load initial-guess document; fall back to bundled JSON or built-in defaults."""
    if path is not None:
        src = Path(path)
        return load_initial_params(src), str(src)
    bundled = bundled_initial_params_path()
    if bundled.is_file():
        return load_initial_params(bundled), str(bundled)
    return build_default_initial_params_document(), "build_default_initial_params_document()"


def _scalar_param_value(value: Any) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _apply_records_to_dataclass(
    base: Any,
    records: Sequence[InitialParamRecord],
    *,
    target_type: str,
) -> Tuple[Any, List[str], List[str]]:
    """Apply scalar records onto a params dataclass; return updated instance."""
    known = {f.name for f in fields(base)}
    applied: List[str] = []
    skipped: List[str] = []
    updates: Dict[str, Any] = {}

    for rec in records:
        field_name, val, _alias = apply_record_with_aliases(
            rec, target_type=target_type, known=known
        )
        if field_name is None or val is None:
            if rec.target_type == target_type:
                skipped.append(rec.name)
            continue
        updates[field_name] = val
        dep_alias = _CHTD_CANONICAL_TO_DEPRECATED_ALIAS.get(field_name)
        if target_type == "chtd" and dep_alias in known:
            updates[dep_alias] = None
        applied.append(field_name)

    if not updates:
        return base, applied, skipped
    return replace(base, **updates), applied, skipped


@dataclass(frozen=True)
class RuntimeParamsBundle:
    """AFE/CHTD model parameters loaded from an initial calibration table."""

    afe_params: AFEParams
    chtd_params: CHTDParams
    initial_params_source: str
    classification: str
    applied_afe: Tuple[str, ...] = ()
    applied_chtd: Tuple[str, ...] = ()
    skipped_names: Tuple[str, ...] = ()

    def provenance_dict(self) -> Dict[str, Any]:
        skipped = list(self.skipped_names)
        return {
            "initial_params_source": self.initial_params_source,
            "classification": self.classification,
            "applied_afe_fields": list(self.applied_afe),
            "applied_chtd_fields": list(self.applied_chtd),
            "applied_afe_count": len(self.applied_afe),
            "applied_chtd_count": len(self.applied_chtd),
            "skipped_param_names": skipped,
            "skipped_count": len(skipped),
            "skipped_preview": skipped[:20],
            "afe_EvaRessCo": float(self.afe_params.EvaRessCo),
            "chtd_CHTD_Dt_P": float(self.chtd_params.CHTD_Dt_P),
        }


def runtime_params_from_document(
    doc: InitialParamsDocument,
    *,
    source: str,
) -> RuntimeParamsBundle:
    """Build AFE/CHTD params from a loaded initial-params document."""
    afe, applied_afe, skipped_afe = _apply_records_to_dataclass(
        AFEParams(), doc.parameters, target_type="airflow"
    )
    chtd, applied_chtd, skipped_chtd = _apply_records_to_dataclass(
        CHTDParams(), doc.parameters, target_type="chtd"
    )
    skipped = tuple(sorted(set(skipped_afe + skipped_chtd)))
    return RuntimeParamsBundle(
        afe_params=afe,
        chtd_params=chtd,
        initial_params_source=source,
        classification=doc.classification,
        applied_afe=tuple(applied_afe),
        applied_chtd=tuple(applied_chtd),
        skipped_names=skipped,
    )


def load_runtime_params(
    path: Optional[Union[str, Path]] = None,
) -> RuntimeParamsBundle:
    """Load AFE/CHTD params for runtime PMV and calibration initial guesses."""
    doc, source = resolve_initial_params(path)
    return runtime_params_from_document(doc, source=source)
