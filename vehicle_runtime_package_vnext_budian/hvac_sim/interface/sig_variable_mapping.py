"""SIG/FTE ↔ PMV interface mapping and workbook table builders."""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

WORKBOOK_VERSION = "1.0.0"
WORKBOOK_SCHEMA = "pmv_sig_interface_workbook_v1"

_CSV_COLUMNS_ZH = (
    "模块",
    "输出层级",
    "类别",
    "变量名",
    "数据类型",
    "描述",
    "信号来源",
    "所属子系统",
)

_SIG_FILTERED_COLUMNS = (
    "module",
    "output_level",
    "category",
    "variable_name",
    "data_type",
    "description",
    "signal_source",
    "subsystem",
    "pmv_relevance",
    "suggested_use",
    "notes",
)

_RUNTIME_INPUT_COLUMNS = (
    "pmv_json_key",
    "pmv_internal_name",
    "required_level",
    "matched_sig_variable",
    "matched_fte_variable",
    "source_description",
    "unit",
    "conversion",
    "available_status",
    "fallback",
    "owner",
    "discussion_needed",
    "priority",
)

_RUNTIME_OUTPUT_COLUMNS = (
    "output_variable",
    "legacy_readable_name",
    "naming_standard_name",
    "data_type",
    "unit",
    "description_cn",
    "source_module",
    "update_rate",
    "valid_range",
    "fallback_behavior",
    "consumer",
    "notes",
)

_NAMING_DICTIONARY_COLUMNS = (
    "token",
    "type",
    "meaning_cn",
    "from_standard",
    "needs_sig_confirmation",
    "examples",
)

_CALIB_DATA_COLUMNS = (
    "calibration_target",
    "data_name",
    "description_cn",
    "sensor_or_tool",
    "unit",
    "sampling_rate",
    "needed_for",
    "required_or_optional",
    "current_status",
    "can_derive_from_sig_output",
    "priority",
    "notes",
)

_CALIB_PARAM_COLUMNS = (
    "parameter_name",
    "description_cn",
    "module",
    "expected_source",
    "current_source",
    "current_value_status",
    "can_measure_directly",
    "needs_pso",
    "priority",
    "owner",
    "notes",
)

_OWNERSHIP_COLUMNS = (
    "variable_prefix",
    "owner_module",
    "created_by",
    "meaning",
    "examples",
    "rules",
)

_OPEN_ISSUES_COLUMNS = (
    "issue_id",
    "topic",
    "description",
    "impact",
    "owner",
    "priority",
    "next_action",
    "status",
)

_HVAC_USEFUL_COLUMNS = (
    "variable_name",
    "module",
    "category",
    "pmv_use_class",
    "description",
    "pmv_json_or_output",
    "notes",
)

_PMV_KEYWORD_PATTERNS: Tuple[Tuple[str, str], ...] = (
    (r"temp|温度|dew|露点", "thermal_boundary"),
    (r"hum|湿度|rh", "humidity"),
    (r"solar|日照|太阳|inten", "solar"),
    (r"spd|speed|车速|veh", "vehicle_motion"),
    (r"eva|蒸发|evap", "hvac_thermal"),
    (r"ptc|加热|hct|heat", "heater"),
    (r"vent|出风|tma|face|foot|吹", "duct_temperature"),
    (r"flow|风量|blow|blwr|airflow", "airflow"),
    (r"defrost|除霜|defog", "defrost"),
    (r"mode|模式|flap|风门|ventila|posn", "hvac_actuator"),
    (r"amb|环境|outside|车外", "ambient"),
    (r"incar|车内|cab|cabin", "cabin_sensor"),
    (r"windshield|玻璃|win", "glass"),
    (r"pmv|comfort|舒适", "comfort"),
)

_MANUAL_SIG_PMV_VARS = frozenset(
    {
        "SIG_AmbTempFb",
        "SIG_AmbTempEn",
        "SIG_AmbTempCorrectedFb",
        "SIG_AmbRelHumFb",
        "SIG_AmbLeSolarIntenFb",
        "SIG_AmbRiSolarIntenFb",
        "SIG_AmbSolarIntenFb",
        "SIG_AmbWindshieldTempFb",
        "SIG_AmbDewPointTempFb",
        "SIG_IpbVehSpdFb",
        "SIG_IpbVehSpdEn",
        "SIG_AcFrntEvaTempFb",
        "SIG_AcFrntIcTempFb",
        "SIG_AcFrntIcTempEn",
        "SIG_AcIcThsTempFb",
        "SIG_AcIcThsRelHumFb",
        "SIG_AcFdDrvFdvTmaFb",
        "SIG_AcFdDrvFdfTmaFb",
        "SIG_AcFpPasFdvTmaFb",
        "SIG_AcFpPasFdfTmaFb",
        "SIG_AcFrntFdDefPosn",
        "SIG_AcFrntModePosn",
        "SIG_AcFrntBlwrFlowFb",
        "SIG_AcFrntAirFlowTar",
        "SIG_PtcFrntOutlTempFb",
        "SIG_PtcFrntInlTempFb",
    }
)

_RUNTIME_INPUT_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "pmv_json_key": "amb_t_c",
        "pmv_internal_name": "RuntimeSignalInput.amb_t_c",
        "required_level": "MUST",
        "matched_sig_variable": "SIG_AmbTempFb",
        "matched_fte_variable": "FTE_AmbTempFb.mean_30s",
        "source_description": "环境温度（VIU/RSM 经 SIG 工程量）",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "none",
        "owner": "SIG/FTE",
        "discussion_needed": "N",
        "priority": "P0",
    },
    {
        "pmv_json_key": "raw_amb_t_c",
        "pmv_internal_name": "u[RawAmbT]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AmbTempFb",
        "matched_fte_variable": "FTE_AmbTempFb.mean_30s",
        "source_description": "未修正环境温度；可与 SIG_AmbTempCorrectedFb 区分",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "amb_t_c",
        "owner": "SIG",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "rh_percent",
        "pmv_internal_name": "PipelineInputs.rh",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AmbRelHumFb",
        "matched_fte_variable": "FTE_AmbRelHumFb.mean_30s",
        "source_description": "相对湿度；备选 SIG_AcIcThsRelHumFb（车内 THS）",
        "unit": "%",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "50% default",
        "owner": "SIG/FTE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "vehicle_speed_kph",
        "pmv_internal_name": "u[VehSpd]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_IpbVehSpdFb",
        "matched_fte_variable": "FTE_IpbVehSpdFb.mean_30s",
        "source_description": "车速 km/h",
        "unit": "km/h",
        "conversion": "direct (validate vs m/s CAN scaling)",
        "available_status": "available_in_sig",
        "fallback": "0 parked",
        "owner": "SIG/FTE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "solar_driver_w_m2",
        "pmv_internal_name": "u[SolarFd]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AmbLeSolarIntenFb",
        "matched_fte_variable": "FTE_AmbLeSolarIntenFb.mean_30s",
        "source_description": "主驾侧/左日照 W/m²",
        "unit": "W/m2",
        "conversion": "validate unit vs RSM raw",
        "available_status": "available_in_sig",
        "fallback": "0 or SIG_AmbSolarIntenFb duplicated",
        "owner": "SIG/FTE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "solar_passenger_w_m2",
        "pmv_internal_name": "u[SolarFp]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AmbRiSolarIntenFb",
        "matched_fte_variable": "FTE_AmbRiSolarIntenFb.mean_30s",
        "source_description": "副驾侧/右日照 W/m²",
        "unit": "W/m2",
        "conversion": "validate unit vs RSM raw",
        "available_status": "available_in_sig",
        "fallback": "0 or solar_w_m2 single sensor",
        "owner": "SIG/FTE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "eva_t_c",
        "pmv_internal_name": "eva_t fallback chain",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AcFrntEvaTempFb",
        "matched_fte_variable": "FTE_AcFrntEvaTempFb.mean_30s",
        "source_description": "前蒸发器温度",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "amb_t_c",
        "owner": "SIG/AFE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "hct_c",
        "pmv_internal_name": "estimate_tma_def hct",
        "required_level": "Optional",
        "matched_sig_variable": "SIG_PtcFrntOutlTempFb",
        "matched_fte_variable": "",
        "source_description": "加热芯出口温度（除霜混合）",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "eva_t precision fallback",
        "owner": "SIG/PTC",
        "discussion_needed": "N",
        "priority": "P2",
    },
    {
        "pmv_json_key": "posn_fdh",
        "pmv_internal_name": "estimate_tma_def posn",
        "required_level": "Optional",
        "matched_sig_variable": "SIG_AcFrntFdDefPosn",
        "matched_fte_variable": "",
        "source_description": "除霜风门位置 0-1",
        "unit": "normalized",
        "conversion": "motor position → normalized",
        "available_status": "available_in_sig",
        "fallback": "eva_t only",
        "owner": "SIG/AFE",
        "discussion_needed": "Y",
        "priority": "P2",
    },
    {
        "pmv_json_key": "windshield_glass_temp_c",
        "pmv_internal_name": "glass observer init",
        "required_level": "Optional",
        "matched_sig_variable": "SIG_AmbWindshieldTempFb",
        "matched_fte_variable": "FTE_AmbWindshieldTempFb.mean_30s",
        "source_description": "前挡玻璃温度观测/初值",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "WinShdTEst observer",
        "owner": "SIG/RSM",
        "discussion_needed": "N",
        "priority": "P2",
    },
    {
        "pmv_json_key": "tma.FrntFdvTma",
        "pmv_internal_name": "u[FrntFdvTma]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AcFdDrvFdvTmaFb",
        "matched_fte_variable": "FTE_AcFdDrvFdvTmaFb.mean_30s",
        "source_description": "主驾吹面出风温度",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "eva_t_c fill all Tma",
        "owner": "SIG/AFE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "tma.FrntFdfTma",
        "pmv_internal_name": "u[FrntFdfTma]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AcFdDrvFdfTmaFb",
        "matched_fte_variable": "FTE_AcFdDrvFdfTmaFb.mean_30s",
        "source_description": "主驾吹脚出风温度",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "eva_t_c",
        "owner": "SIG/AFE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "tma.FrntFpvTma",
        "pmv_internal_name": "u[FrntFpvTma]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AcFpPasFdvTmaFb",
        "matched_fte_variable": "FTE_AcFpPasFdvTmaFb.mean_30s",
        "source_description": "副驾吹面出风温度",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "eva_t_c",
        "owner": "SIG/AFE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "tma.FrntFpfTma",
        "pmv_internal_name": "u[FrntFpfTma]",
        "required_level": "Recommended",
        "matched_sig_variable": "SIG_AcFpPasFdfTmaFb",
        "matched_fte_variable": "FTE_AcFpPasFdfTmaFb.mean_30s",
        "source_description": "副驾吹脚出风温度",
        "unit": "degC",
        "conversion": "direct",
        "available_status": "available_in_sig",
        "fallback": "eva_t_c",
        "owner": "SIG/AFE",
        "discussion_needed": "N",
        "priority": "P1",
    },
    {
        "pmv_json_key": "driver_face_flow",
        "pmv_internal_name": "u[FrntFdvFlow]",
        "required_level": "Recommended",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "主驾吹面风量 m³/h；来自 AFE_Calc 或 bench 显式注入",
        "unit": "m3/h",
        "conversion": "AFE Q_out or explicit runtime",
        "available_status": "needs_afe_or_explicit",
        "fallback": "0 u-bus",
        "owner": "AFE/FTE",
        "discussion_needed": "Y",
        "priority": "P0",
    },
    {
        "pmv_json_key": "passenger_face_flow",
        "pmv_internal_name": "u[FrntFpvFlow]",
        "required_level": "Recommended",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "副驾吹面风量",
        "unit": "m3/h",
        "conversion": "AFE Q_out or explicit",
        "available_status": "needs_afe_or_explicit",
        "fallback": "0",
        "owner": "AFE/FTE",
        "discussion_needed": "Y",
        "priority": "P0",
    },
    {
        "pmv_json_key": "driver_floor_flow",
        "pmv_internal_name": "u[FrntFdfFlow]",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "主驾吹脚风量",
        "unit": "m3/h",
        "conversion": "AFE Q_out",
        "available_status": "needs_afe_or_explicit",
        "fallback": "0",
        "owner": "AFE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "driver_defrost_flow",
        "pmv_internal_name": "u[FrntFdDefFlow]",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "主驾除霜风量",
        "unit": "m3/h",
        "conversion": "AFE Q_out",
        "available_status": "needs_afe_or_explicit",
        "fallback": "0",
        "owner": "AFE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "image_inputs.driver.occupied",
        "pmv_internal_name": "SeatOccupantInput.occupied",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "图像/DMS 占用检测",
        "unit": "bool",
        "conversion": "image module v1",
        "available_status": "external_image_module",
        "fallback": "default true",
        "owner": "Image/FTE",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "image_inputs.driver.ir_head_surface_temp_c",
        "pmv_internal_name": "IR surface temp",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "IR 头部表面温度（非 CHTD 状态）",
        "unit": "degC",
        "conversion": "IR module",
        "available_status": "external_image_module",
        "fallback": "model_only PMV air temp",
        "owner": "Image/IR",
        "discussion_needed": "Y",
        "priority": "P1",
    },
    {
        "pmv_json_key": "image_inputs.driver.confidence",
        "pmv_internal_name": "IR confidence gate",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "IR/占用置信度",
        "unit": "0-1",
        "conversion": "direct",
        "available_status": "external_image_module",
        "fallback": "fusion skipped",
        "owner": "Image",
        "discussion_needed": "N",
        "priority": "P2",
    },
    {
        "pmv_json_key": "image_inputs.driver.activity_met",
        "pmv_internal_name": "met_driver",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "代谢率 MET",
        "unit": "met",
        "conversion": "direct or class map",
        "available_status": "external_image_module",
        "fallback": "1.0",
        "owner": "Image/FTE",
        "discussion_needed": "N",
        "priority": "P2",
    },
    {
        "pmv_json_key": "image_inputs.driver.clothing_clo",
        "pmv_internal_name": "clo_driver",
        "required_level": "Optional",
        "matched_sig_variable": "",
        "matched_fte_variable": "",
        "source_description": "服装热阻 clo",
        "unit": "clo",
        "conversion": "class → clo table",
        "available_status": "external_image_module",
        "fallback": "0.5",
        "owner": "Image/FTE",
        "discussion_needed": "N",
        "priority": "P2",
    },
    {
        "pmv_json_key": "ac_mode_ventila_posn",
        "pmv_internal_name": "FlapSet decode (future)",
        "required_level": "Optional",
        "matched_sig_variable": "SIG_AcFrntModePosn",
        "matched_fte_variable": "",
        "source_description": "模式风门枚举；AFE-I01 待 decode",
        "unit": "enum",
        "conversion": "HVAC mode table → flap",
        "available_status": "needs_decode_module",
        "fallback": "default flaps",
        "owner": "SIG/AFE",
        "discussion_needed": "Y",
        "priority": "P0",
    },
)

def _runtime_output_row(
    *,
    naming_standard_name: str,
    legacy_readable_name: str,
    data_type: str,
    unit: str,
    description_cn: str,
    source_module: str,
    update_rate: str = "1 Hz",
    valid_range: str = "",
    fallback_behavior: str = "",
    consumer: str = "",
    notes: str = "",
) -> Dict[str, Any]:
    """Build one Runtime_Output_PMv row; ``output_variable`` = standard SIG-style name."""
    return {
        "output_variable": naming_standard_name,
        "legacy_readable_name": legacy_readable_name,
        "naming_standard_name": naming_standard_name,
        "data_type": data_type,
        "unit": unit,
        "description_cn": description_cn,
        "source_module": source_module,
        "update_rate": update_rate,
        "valid_range": valid_range,
        "fallback_behavior": fallback_behavior,
        "consumer": consumer,
        "notes": notes,
    }


_RUNTIME_OUTPUT_ROWS: Tuple[Dict[str, Any], ...] = (
    _runtime_output_row(
        naming_standard_name="PMV_FdPmvFb",
        legacy_readable_name="PMV_DriverPmv",
        data_type="float",
        unit="-",
        description_cn="主驾 PMV 基准值",
        source_module="pmv/Fanger",
        valid_range="-3~+3",
        fallback_behavior="finite via degradation",
        consumer="HMI/CDC",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpPmvFb",
        legacy_readable_name="PMV_PassengerPmv",
        data_type="float",
        unit="-",
        description_cn="副驾 PMV 基准值",
        source_module="pmv/Fanger",
        valid_range="-3~+3",
        fallback_behavior="finite",
        consumer="HMI/CDC",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FdPpdFb",
        legacy_readable_name="PMV_DriverPpd",
        data_type="float",
        unit="%",
        description_cn="主驾 PPD",
        source_module="pmv/Fanger",
        valid_range="0~100",
        fallback_behavior="finite",
        consumer="HMI",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpPpdFb",
        legacy_readable_name="PMV_PassengerPpd",
        data_type="float",
        unit="%",
        description_cn="副驾 PPD",
        source_module="pmv/Fanger",
        valid_range="0~100",
        fallback_behavior="finite",
        consumer="HMI",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FdPmvEn",
        legacy_readable_name="PMV_DriverValid",
        data_type="bool",
        unit="-",
        description_cn="主驾 PMV 结果有效（occupied）",
        source_module="occupant",
        valid_range="true/false",
        fallback_behavior="false if unoccupied",
        consumer="HMI",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpPmvEn",
        legacy_readable_name="PMV_PassengerValid",
        data_type="bool",
        unit="-",
        description_cn="副驾 PMV 结果有效",
        source_module="occupant",
        valid_range="true/false",
        fallback_behavior="false if unoccupied",
        consumer="HMI",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FdIcTempFb",
        legacy_readable_name="PMV_DriverAirTemp",
        data_type="float",
        unit="degC",
        description_cn="主驾 PMV 用空气温度",
        source_module="pipeline/CHTD+IR",
        valid_range="-40~90",
        fallback_behavior="model or fused",
        consumer="internal",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpIcTempFb",
        legacy_readable_name="PMV_PassengerAirTemp",
        data_type="float",
        unit="degC",
        description_cn="副驾 PMV 用空气温度",
        source_module="pipeline",
        valid_range="-40~90",
        fallback_behavior="model or fused",
        consumer="internal",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FdMrtFb",
        legacy_readable_name="PMV_DriverMrt",
        data_type="float",
        unit="degC",
        description_cn="主驾平均辐射温度",
        source_module="CHTD",
        valid_range="-40~90",
        fallback_behavior="zone MRT proxy",
        consumer="internal",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpMrtFb",
        legacy_readable_name="PMV_PassengerMrt",
        data_type="float",
        unit="degC",
        description_cn="副驾平均辐射温度",
        source_module="CHTD",
        valid_range="-40~90",
        fallback_behavior="zone MRT proxy",
        consumer="internal",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FdAirSpdFb",
        legacy_readable_name="PMV_DriverAirSpeed",
        data_type="float",
        unit="m/s",
        description_cn="主驾头部局部风速",
        source_module="air_speed",
        valid_range="0~2",
        fallback_behavior="geometry placeholder",
        consumer="PMV",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FpAirSpdFb",
        legacy_readable_name="PMV_PassengerAirSpeed",
        data_type="float",
        unit="m/s",
        description_cn="副驾头部局部风速",
        source_module="air_speed",
        valid_range="0~2",
        fallback_behavior="geometry placeholder",
        consumer="PMV",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_QualityIndcr",
        legacy_readable_name="PMV_QualityLevel",
        data_type="string",
        unit="-",
        description_cn="PMV 质量等级 full/degraded/baseline_only",
        source_module="runtime_diagnostics",
        valid_range="enum",
        fallback_behavior="always emitted",
        consumer="HMI/diag",
        notes="Indcr=indicator；待 SIG 确认枚举编码",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_IrFusionModeFb",
        legacy_readable_name="PMV_IrFusionStatus",
        data_type="string",
        unit="-",
        description_cn="IR 融合状态 used/model_only/disabled/missing",
        source_module="runtime_diagnostics",
        valid_range="enum",
        fallback_behavior="default disabled",
        consumer="diag",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_FallbackIndcr",
        legacy_readable_name="PMV_FallbackUsed",
        data_type="object",
        unit="-",
        description_cn="回退路径摘要 JSON",
        source_module="runtime_diagnostics",
        fallback_behavior="additive",
        consumer="diag",
    ),
    _runtime_output_row(
        naming_standard_name="PMV_ChtdModeFb",
        legacy_readable_name="PMV_ChtdStatus",
        data_type="string",
        unit="-",
        description_cn="CHTD trace 状态摘要",
        source_module="pipeline",
        fallback_behavior="AS_FOUND default",
        consumer="diag",
    ),
)

_PMV_NAMING_DICTIONARY_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "token": "PMV",
        "type": "module",
        "meaning_cn": "PMV 舒适度模块前缀",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "PMV_FdPmvFb",
    },
    {
        "token": "Fd",
        "type": "scope",
        "meaning_cn": "前排主驾（Front Driver）",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcFdDrvTempFb, PMV_FdPmvFb",
    },
    {
        "token": "Fp",
        "type": "scope",
        "meaning_cn": "前排副驾（Front Passenger）",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcFpPasTempFb, PMV_FpPmvFb",
    },
    {
        "token": "Ic",
        "type": "quantity",
        "meaning_cn": "车内（In-Car）空气量",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcFrntIcTempFb, PMV_FdIcTempFb",
    },
    {
        "token": "AirSpd",
        "type": "quantity",
        "meaning_cn": "局部空气流速",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "PMV_FdAirSpdFb",
    },
    {
        "token": "Pmv",
        "type": "quantity",
        "meaning_cn": "Predicted Mean Vote 热舒适指标",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_FdPmvFb",
    },
    {
        "token": "Ppd",
        "type": "quantity",
        "meaning_cn": "Predicted Percentage Dissatisfied",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_FdPpdFb",
    },
    {
        "token": "Mrt",
        "type": "quantity",
        "meaning_cn": "Mean Radiant Temperature 平均辐射温度",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_FdMrtFb",
    },
    {
        "token": "Quality",
        "type": "quantity",
        "meaning_cn": "输出质量/可信度等级",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_QualityIndcr",
    },
    {
        "token": "Fallback",
        "type": "quantity",
        "meaning_cn": "降级/回退路径指示",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_FallbackIndcr",
    },
    {
        "token": "Fusion",
        "type": "quantity",
        "meaning_cn": "IR 与模型融合模式",
        "from_standard": "no",
        "needs_sig_confirmation": "yes",
        "examples": "PMV_IrFusionModeFb",
    },
    {
        "token": "Fb",
        "type": "action_type",
        "meaning_cn": "Feedback 反馈量（实测/计算输出）",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcFdDrvFdvTmaFb, PMV_FdPmvFb",
    },
    {
        "token": "En",
        "type": "action_type",
        "meaning_cn": "Enable/Valid 有效标志",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcFrntIcTempEn, PMV_FdPmvEn",
    },
    {
        "token": "Indcr",
        "type": "action_type",
        "meaning_cn": "Indicator 指示量/枚举编码",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "SIG_AcAmbPm25Indcr, PMV_QualityIndcr",
    },
    {
        "token": "Mode",
        "type": "action_type",
        "meaning_cn": "模式/状态字",
        "from_standard": "yes",
        "needs_sig_confirmation": "no",
        "examples": "PMV_ChtdModeFb, PMV_IrFusionModeFb",
    },
)

_CALIBRATION_DATA_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "calibration_target": "AFE L3",
        "data_name": "bench_measured_vent_flow_m3h",
        "description_cn": "台架实测各风口体积流量",
        "sensor_or_tool": "风量计/风洞",
        "unit": "m3/h",
        "sampling_rate": "1 Hz",
        "needed_for": "AFE PSO resistance/fan",
        "required_or_optional": "Required",
        "current_status": "schema_ready",
        "can_derive_from_sig_output": "N",
        "priority": "P0",
        "notes": "非 runtime 输入",
    },
    {
        "calibration_target": "air_speed L2",
        "data_name": "outlet_effective_area_m2",
        "description_cn": "各出风口等效面积",
        "sensor_or_tool": "PIV/热线/台架",
        "unit": "m2",
        "sampling_rate": "once per vehicle",
        "needed_for": "local air speed / PMV vel",
        "required_or_optional": "Required",
        "current_status": "geometry placeholder",
        "can_derive_from_sig_output": "N",
        "priority": "P0",
        "notes": "见 geometry_defaults.json",
    },
    {
        "calibration_target": "air_speed L1",
        "data_name": "head_probe_wind_speed_m_s",
        "description_cn": "头部探针点风速",
        "sensor_or_tool": "风速仪",
        "unit": "m/s",
        "sampling_rate": "1-10 Hz",
        "needed_for": "validate air_speed.py",
        "required_or_optional": "Recommended",
        "current_status": "wind_speed CSV schema",
        "can_derive_from_sig_output": "N",
        "priority": "P1",
        "notes": "验证用，非 runtime",
    },
    {
        "calibration_target": "CHTD PSO",
        "data_name": "zone_temperature_time_series",
        "description_cn": "各温区实测温度时间序列",
        "sensor_or_tool": "热电偶/PT100",
        "unit": "degC",
        "sampling_rate": "0.2-1 Hz",
        "needed_for": "CHTD parameter PSO",
        "required_or_optional": "Required",
        "current_status": "required for calibration only",
        "can_derive_from_sig_output": "partial SIG_AcFrntIcTempFb",
        "priority": "P0",
        "notes": "runtime PMV 不要求",
    },
    {
        "calibration_target": "IR/EKF",
        "data_name": "ir_head_surface_temp_c",
        "description_cn": "IR 头部表面温度同步记录",
        "sensor_or_tool": "IR camera",
        "unit": "degC",
        "sampling_rate": "1-5 Hz",
        "needed_for": "IR fusion tuning",
        "required_or_optional": "Optional",
        "current_status": "image_inputs v1",
        "can_derive_from_sig_output": "N",
        "priority": "P1",
        "notes": "图像模块输出",
    },
    {
        "calibration_target": "defrost/glass",
        "data_name": "windshield_glass_temp_c",
        "description_cn": "前挡玻璃/内表面温度",
        "sensor_or_tool": "IR/贴点",
        "unit": "degC",
        "sampling_rate": "0.2 Hz",
        "needed_for": "glass observer validation",
        "required_or_optional": "Optional",
        "current_status": "SIG_AmbWindshieldTempFb available",
        "can_derive_from_sig_output": "Y",
        "priority": "P2",
        "notes": "",
    },
)

_OWNERSHIP_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "variable_prefix": "SIG_",
        "owner_module": "SIG_IO",
        "created_by": "CAN/DBC decode + derived",
        "meaning": "Snapshot 工程量（CAN + 衍生）",
        "examples": "SIG_AmbTempFb, SIG_AcFrntEvaTempFb",
        "rules": "PMV 不解析 CAN bit；只消费 SIG 工程量",
    },
    {
        "variable_prefix": "FTE_",
        "owner_module": "FeatureCore",
        "created_by": "FTE feature engine",
        "meaning": "时序特征（均值/斜率/峰值）",
        "examples": "FTE_AmbTempFb.mean_30s",
        "rules": "可选替代瞬时 SIG；需注明采样窗",
    },
    {
        "variable_prefix": "AFE_",
        "owner_module": "AFE/AFE_Calc",
        "created_by": "Simulink AFE workspace",
        "meaning": "风道阻力/风机/workspace 参数",
        "examples": "AFE_FHEvaRessCo_P",
        "rules": "风量由 AFE_Calc 求解；runtime 可显式 override flow",
    },
    {
        "variable_prefix": "CHTD_",
        "owner_module": "CHTD thermal",
        "created_by": "UpdateParam.m / calibration",
        "meaning": "热模型标量/LUT 参数",
        "examples": "CHTD_CabinFdAmbConvCo_M",
        "rules": "工程师参数；非 SIG 输出",
    },
    {
        "variable_prefix": "PMV_",
        "owner_module": "PMV comfort output",
        "created_by": "run_runtime_pmv / pipeline",
        "meaning": "对外 PMV 舒适度输出",
        "examples": "PMV_DriverPmv, PMV_QualityLevel",
        "rules": "所有对外 comfort 输出必须 PMV_ 前缀",
    },
    {
        "variable_prefix": "PMV_Diag_",
        "owner_module": "PMV diagnostics",
        "created_by": "runtime_diagnostics",
        "meaning": "诊断/降级/追溯",
        "examples": "PMV_Diag_FallbackUsed (future alias)",
        "rules": "不影响 comfort 数值；可映射 degradation_report",
    },
    {
        "variable_prefix": "PMV_Calib_",
        "owner_module": "calibration",
        "created_by": "PSO/bench pipeline",
        "meaning": "标定过程中间量/数据集字段",
        "examples": "PMV_Calib_AfeResidual",
        "rules": "非 runtime 必填",
    },
    {
        "variable_prefix": "PMV_Rec_",
        "owner_module": "recommendation (future)",
        "created_by": "future reco module",
        "meaning": "推荐/策略输出（非 baseline PMV）",
        "examples": "PMV_Rec_TargetTemp",
        "rules": "与 baseline PMV_ 输出分离",
    },
)

_OPEN_ISSUES_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "issue_id": "AFE-I01",
        "topic": "AC mode decode",
        "description": "AC_ModeVentilaPosn 未映射 FlapSet",
        "impact": "吹脚/除霜模式风量错误",
        "owner": "HVAC/SIG",
        "priority": "P0",
        "next_action": "提供 0-15 decode 表 → signal_conditioning",
        "status": "open",
    },
    {
        "issue_id": "AFE-I07",
        "topic": "AFE params vehicle scope",
        "description": "param.xlsx 来自其他车型",
        "impact": "绝对风量/对流量级偏差",
        "owner": "标定",
        "priority": "P0",
        "next_action": "L3 bench PSO",
        "status": "open",
    },
    {
        "issue_id": "CHTD-P0-LUT",
        "topic": "CHTD LUT placeholders",
        "description": "大量 CHTD_*_M 仍为全 1 占位",
        "impact": "热交换系数不准",
        "owner": "热标定",
        "priority": "P0",
        "next_action": "见 parameter_p0_review.csv",
        "status": "open",
    },
    {
        "issue_id": "PMV-FLOW-01",
        "topic": "Runtime vent flow source",
        "description": "SIG 无直接 m³/h 风口流量；需 AFE 或显式注入",
        "impact": "minimal runtime 时 flow=0",
        "owner": "PMV/AFE",
        "priority": "P0",
        "next_action": "对接 AFE Q_out 或 FTE 映射",
        "status": "open",
    },
    {
        "issue_id": "PMV-IR-01",
        "topic": "IR in SIG catalog",
        "description": "IR/occupied 不在 output_variables.csv",
        "impact": "图像模块对接需单独接口",
        "owner": "Image/FTE",
        "priority": "P1",
        "next_action": "补充 image_inputs 对接表",
        "status": "open",
    },
    {
        "issue_id": "PMV-GEO-01",
        "topic": "Geometry effective area",
        "description": "geometry_defaults measured=false",
        "impact": "local air speed / PMV vel",
        "owner": "标定/总布置",
        "priority": "P0",
        "next_action": "台架测 effective area",
        "status": "open",
    },
)


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def default_output_variables_path() -> Path:
    return _repo_root() / "files" / "output_variables.csv"


def default_parameter_master_path() -> Path:
    return (
        _repo_root()
        / "simulink_conversion_package"
        / "python_targets"
        / "parameter_master_table.json"
    )


def load_output_variables_csv(path: Optional[Path] = None) -> List[Dict[str, str]]:
    csv_path = path or default_output_variables_path()
    rows: List[Dict[str, str]] = []
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        for raw in reader:
            row = {
                "module": (raw.get("模块") or "").strip(),
                "output_level": (raw.get("输出层级") or "").strip(),
                "category": (raw.get("类别") or "").strip(),
                "variable_name": (raw.get("变量名") or "").strip(),
                "data_type": (raw.get("数据类型") or "").strip(),
                "description": (raw.get("描述") or "").strip(),
                "signal_source": (raw.get("信号来源") or "").strip(),
                "subsystem": (raw.get("所属子系统") or "").strip(),
            }
            if row["variable_name"]:
                rows.append(row)
    return rows


def _pmv_relevance_for_row(row: Mapping[str, str]) -> Tuple[str, str, str]:
    name = row.get("variable_name", "")
    desc = row.get("description", "")
    text = f"{name} {desc} {row.get('signal_source','')}".lower()
    if name in _MANUAL_SIG_PMV_VARS or name.startswith("FTE_Amb") or name.startswith("FTE_Ac"):
        relevance = "high"
    else:
        relevance = "low"
        matched_tag = ""
        for pattern, tag in _PMV_KEYWORD_PATTERNS:
            if re.search(pattern, text, re.IGNORECASE):
                relevance = "medium" if relevance == "low" else relevance
                matched_tag = tag
                if name.startswith("SIG_") and tag in {
                    "thermal_boundary",
                    "humidity",
                    "solar",
                    "vehicle_motion",
                    "duct_temperature",
                    "airflow",
                    "defrost",
                    "hvac_actuator",
                    "ambient",
                    "cabin_sensor",
                    "glass",
                }:
                    relevance = "high"
                    break
        if relevance == "low" and matched_tag:
            relevance = "medium"

    if name in _MANUAL_SIG_PMV_VARS:
        relevance = "high"

    use_map = {
        "high": "runtime_adapter / CHTD boundary / AFE input",
        "medium": "conditioning / observer / diagnostic",
        "low": "not primary for PMV v1",
    }
    suggested = use_map.get(relevance, "review")
    notes = ""
    if name.startswith("FTE_"):
        notes = "FTE 特征；可用于 runtime conditioning 或慢变输入"
    elif row.get("module") == "SIG_IO" and "derived" in row.get("output_level", ""):
        notes = "SIG 衍生量；优先于原始 CAN"
    return relevance, suggested, notes


def filter_pmv_sig_variables(rows: Sequence[Mapping[str, str]]) -> List[Dict[str, Any]]:
    filtered: List[Dict[str, Any]] = []
    for row in rows:
        relevance, suggested, notes = _pmv_relevance_for_row(row)
        if relevance == "low":
            continue
        filtered.append(
            {
                **dict(row),
                "pmv_relevance": relevance,
                "suggested_use": suggested,
                "notes": notes,
            }
        )
    filtered.sort(key=lambda r: (r["pmv_relevance"] != "high", r["variable_name"]))
    return filtered


def _hvac_use_class(row: Mapping[str, str]) -> str:
    name = row.get("variable_name", "")
    desc = row.get("description", "")
    rel, _, _ = _pmv_relevance_for_row(row)
    if rel == "low":
        return "diagnostic only"
    if "observer" in desc or "玻璃" in desc or "Windshield" in name:
        return "observer input"
    if "flow" in desc.lower() or "风量" in desc or "Blwr" in name:
        return "AFE input"
    if name.startswith("FTE_"):
        return "runtime conditioning input"
    if "Temp" in name and ("Ic" in name or "车内" in desc):
        return "calibration log"
    if rel == "high":
        return "direct model input"
    return "diagnostic only"


def build_hvac_useful_sig_fields(
    filtered_rows: Sequence[Mapping[str, Any]],
) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for row in filtered_rows:
        if row.get("pmv_relevance") not in ("high", "medium"):
            continue
        name = row["variable_name"]
        pmv_key = ""
        for rt in _RUNTIME_INPUT_ROWS:
            if rt.get("matched_sig_variable") == name:
                pmv_key = rt["pmv_json_key"]
                break
        out.append(
            {
                "variable_name": name,
                "module": row.get("module", ""),
                "category": row.get("category", ""),
                "pmv_use_class": _hvac_use_class(row),
                "description": row.get("description", ""),
                "pmv_json_or_output": pmv_key or "see Runtime_Input_Mapping",
                "notes": row.get("notes", ""),
            }
        )
    # ensure key manual vars present even if filtered missed
    existing = {r["variable_name"] for r in out}
    for sig in sorted(_MANUAL_SIG_PMV_VARS):
        if sig not in existing:
            out.append(
                {
                    "variable_name": sig,
                    "module": "SIG_IO",
                    "category": "manual",
                    "pmv_use_class": "direct model input",
                    "description": "manual PMV mapping",
                    "pmv_json_or_output": "see Runtime_Input_Mapping",
                    "notes": "hand-maintained",
                }
            )
    return out


def build_calibration_parameter_request(
    parameter_master_path: Optional[Path] = None,
    *,
    max_rows: int = 80,
) -> List[Dict[str, Any]]:
    path = parameter_master_path or default_parameter_master_path()
    rows: List[Dict[str, Any]] = []
    if not path.is_file():
        return _fallback_calibration_params()

    doc = json.loads(path.read_text(encoding="utf-8"))
    params = doc.get("parameters", [])
    # P0 first, then PSO candidates
    ordered = sorted(
        params,
        key=lambda p: (
            p.get("priority") != "P0",
            not p.get("pso_candidate"),
            p.get("parameter_name", ""),
        ),
    )
    for p in ordered[:max_rows]:
        status = "placeholder" if p.get("is_placeholder_like") else "calibrated_or_default"
        if p.get("source_status") == "missing_in_excel":
            status = "missing_in_excel"
        rows.append(
            {
                "parameter_name": p.get("parameter_name", ""),
                "description_cn": p.get("chinese_description", ""),
                "module": p.get("module", ""),
                "expected_source": "vehicle calibration / param.xlsx",
                "current_source": p.get("source_file", ""),
                "current_value_status": status,
                "can_measure_directly": "Y" if p.get("module") == "geometry" else "N",
                "needs_pso": "Y" if p.get("pso_candidate") else "N",
                "priority": p.get("priority", "P2"),
                "owner": "热标定" if p.get("module") == "CHTD" else "AFE标定",
                "notes": f"LUT={p.get('is_lut')} placeholder_like={p.get('is_placeholder_like')}",
            }
        )
    return rows


def _fallback_calibration_params() -> List[Dict[str, Any]]:
    return [
        {
            "parameter_name": "EvaRessCo",
            "description_cn": "流阻系数",
            "module": "AFE",
            "expected_source": "param.xlsx",
            "current_source": "python default",
            "current_value_status": "from_other_vehicle",
            "can_measure_directly": "N",
            "needs_pso": "Y",
            "priority": "P0",
            "owner": "AFE标定",
            "notes": "parameter_master_table.json missing",
        }
    ]


def build_readme_rows(
    *,
    generated_at: str,
    input_files: Sequence[str],
) -> List[List[str]]:
    return [
        ["PMV SIG Interface Workbook"],
        ["版本", WORKBOOK_VERSION],
        ["生成时间 UTC", generated_at],
        ["Schema", WORKBOOK_SCHEMA],
        [""],
        ["输入文件"],
        *[[f, ""] for f in input_files],
        [""],
        ["三类数据说明（勿混用）"],
        ["类型", "含义", "示例"],
        ["Runtime 数据", "每步 PMV 运行所需 SIG/FTE/AFE/图像工程量", "amb_t_c, driver_face_flow"],
        ["标定采集数据", "台架/路试实验记录，用于 PSO 与验证", "bench Qm, 探针风速"],
        ["工程师参数", "模型系数/LUT/几何，写入 params", "CHTD_*_M, AFE_FH*"],
        [""],
        ["命名规则", "模块名_作用域作用量作用类型_标定数据类型；见 PMV_Naming_Dictionary"],
        ["维护", "python_impl/hvac_sim/interface/sig_variable_mapping.py"],
        ["生成", "python_impl/examples/generate_pmv_interface_workbook.py"],
    ]


@dataclass
class WorkbookBuildResult:
    xlsx_path: Optional[Path]
    csv_dir: Path
    sheet_names: List[str]
    sig_filtered_count: int
    sig_matched_runtime_count: int
    fte_matched_runtime_count: int
    runtime_missing_fields: List[str]
    discussion_fields: List[str]
    pmv_output_count: int


def _write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in columns})


def _style_workbook_sheet(ws, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    p0_fill = PatternFill("solid", fgColor="FFC7CE")
    p1_fill = PatternFill("solid", fgColor="FFEB9C")

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = max(12, min(40, len(col_name) + 8))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, row in enumerate(rows, start=2):
        priority = str(row.get("priority", ""))
        for col_idx, col_name in enumerate(columns, start=1):
            value = row.get(col_name, "")
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if priority == "P0":
                cell.fill = p0_fill
            elif priority == "P1":
                cell.fill = p1_fill

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def build_workbook_tables(
    *,
    output_variables_path: Optional[Path] = None,
    parameter_master_path: Optional[Path] = None,
) -> Dict[str, Any]:
    all_rows = load_output_variables_csv(output_variables_path)
    sig_filtered = filter_pmv_sig_variables(all_rows)
    runtime_input = [dict(r) for r in _RUNTIME_INPUT_ROWS]
    runtime_output = [dict(r) for r in _RUNTIME_OUTPUT_ROWS]
    naming_dictionary = [dict(r) for r in _PMV_NAMING_DICTIONARY_ROWS]
    calib_data = [dict(r) for r in _CALIBRATION_DATA_ROWS]
    calib_params = build_calibration_parameter_request(parameter_master_path)
    ownership = [dict(r) for r in _OWNERSHIP_ROWS]
    open_issues = [dict(r) for r in _OPEN_ISSUES_ROWS]
    hvac_useful = build_hvac_useful_sig_fields(sig_filtered)

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    input_files = [
        str(output_variables_path or default_output_variables_path()),
        str(parameter_master_path or default_parameter_master_path()),
        "files/信号命名规范(1).xlsx",
        "python_impl/hvac_sim/runtime_adapter.py",
        "simulink_conversion_package/python_targets/RUNTIME_DEGRADATION_CODE_STATUS.md",
    ]
    readme = build_readme_rows(generated_at=generated_at, input_files=input_files)

    runtime_missing = [
        r["pmv_json_key"]
        for r in runtime_input
        if r.get("available_status")
        in ("needs_afe_or_explicit", "needs_decode_module", "external_image_module")
    ]
    discussion = [
        r["pmv_json_key"] for r in runtime_input if str(r.get("discussion_needed", "")).upper() == "Y"
    ]
    sig_matched = sum(1 for r in runtime_input if r.get("matched_sig_variable"))
    fte_matched = sum(1 for r in runtime_input if r.get("matched_fte_variable"))

    return {
        "readme": readme,
        "SIG_Output_Variables_Filtered": sig_filtered,
        "Runtime_Input_Mapping": runtime_input,
        "Runtime_Output_PMv": runtime_output,
        "PMV_Naming_Dictionary": naming_dictionary,
        "Calibration_Data_To_Collect": calib_data,
        "Calibration_Parameter_Request": calib_params,
        "Variable_Ownership_Naming": ownership,
        "Open_Issues": open_issues,
        "HVAC_Model_Useful_SIG_Fields": hvac_useful,
        "meta": {
            "generated_at": generated_at,
            "sig_filtered_count": len(sig_filtered),
            "sig_matched_runtime_count": sig_matched,
            "fte_matched_runtime_count": fte_matched,
            "runtime_missing_fields": runtime_missing,
            "discussion_fields": discussion,
            "pmv_output_count": len(runtime_output),
        },
    }


def write_pmv_interface_workbook(
    output_xlsx: Path,
    csv_dir: Path,
    *,
    output_variables_path: Optional[Path] = None,
    parameter_master_path: Optional[Path] = None,
) -> WorkbookBuildResult:
    tables = build_workbook_tables(
        output_variables_path=output_variables_path,
        parameter_master_path=parameter_master_path,
    )
    meta = tables["meta"]
    csv_dir.mkdir(parents=True, exist_ok=True)

    sheet_specs: List[Tuple[str, Sequence[str], Any]] = [
        ("README", ("A", "B", "C"), tables["readme"]),
        ("SIG_Output_Variables_Filtered", _SIG_FILTERED_COLUMNS, tables["SIG_Output_Variables_Filtered"]),
        ("Runtime_Input_Mapping", _RUNTIME_INPUT_COLUMNS, tables["Runtime_Input_Mapping"]),
        ("Runtime_Output_PMv", _RUNTIME_OUTPUT_COLUMNS, tables["Runtime_Output_PMv"]),
        ("PMV_Naming_Dictionary", _NAMING_DICTIONARY_COLUMNS, tables["PMV_Naming_Dictionary"]),
        ("Calibration_Data_To_Collect", _CALIB_DATA_COLUMNS, tables["Calibration_Data_To_Collect"]),
        ("Calibration_Parameter_Request", _CALIB_PARAM_COLUMNS, tables["Calibration_Parameter_Request"]),
        ("Variable_Ownership_Naming", _OWNERSHIP_COLUMNS, tables["Variable_Ownership_Naming"]),
        ("Open_Issues", _OPEN_ISSUES_COLUMNS, tables["Open_Issues"]),
        ("HVAC_Model_Useful_SIG_Fields", _HVAC_USEFUL_COLUMNS, tables["HVAC_Model_Useful_SIG_Fields"]),
    ]

    for sheet_name, columns, data in sheet_specs[1:]:
        _write_csv(csv_dir / f"{sheet_name}.csv", columns, data)

    xlsx_path: Optional[Path] = None
    sheet_names = [s[0] for s in sheet_specs]
    try:
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)

        for sheet_name, columns, data in sheet_specs:
            ws = wb.create_sheet(title=sheet_name[:31])
            if sheet_name == "README":
                for r_idx, row in enumerate(data, start=1):
                    for c_idx, val in enumerate(row, start=1):
                        ws.cell(row=r_idx, column=c_idx, value=val)
                ws.freeze_panes = "A2"
                ws.column_dimensions["A"].width = 28
                ws.column_dimensions["B"].width = 48
                ws.column_dimensions["C"].width = 48
            else:
                for c_idx, col in enumerate(columns, start=1):
                    ws.cell(row=1, column=c_idx, value=col)
                for r_idx, row in enumerate(data, start=2):
                    for c_idx, col in enumerate(columns, start=1):
                        val = row.get(col, "")
                        if isinstance(val, (list, dict)):
                            val = json.dumps(val, ensure_ascii=False)
                        ws.cell(row=r_idx, column=c_idx, value=val)
                _style_workbook_sheet(ws, columns, data)

        output_xlsx.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_xlsx)
        xlsx_path = output_xlsx
    except ImportError:
        xlsx_path = None

    return WorkbookBuildResult(
        xlsx_path=xlsx_path,
        csv_dir=csv_dir,
        sheet_names=sheet_names,
        sig_filtered_count=int(meta["sig_filtered_count"]),
        sig_matched_runtime_count=int(meta["sig_matched_runtime_count"]),
        fte_matched_runtime_count=int(meta["fte_matched_runtime_count"]),
        runtime_missing_fields=list(meta["runtime_missing_fields"]),
        discussion_fields=list(meta["discussion_fields"]),
        pmv_output_count=int(meta["pmv_output_count"]),
    )
